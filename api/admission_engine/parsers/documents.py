import csv
import io
import zipfile
from dataclasses import dataclass
from pathlib import Path

import pymupdf
from docx import Document

from api.admission_engine.errors import DomainError

MAX_UPLOAD_BYTES = 4 * 1024 * 1024
MAX_EXTRACTED_CHARACTERS = 100_000
MAX_ARCHIVE_FILES = 2_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_DOCX_XML_BYTES = 5 * 1024 * 1024
MAX_PDF_PAGES = 200
ALLOWED_EXTENSIONS = frozenset({".txt", ".md", ".docx", ".pdf", ".csv"})


class UnsupportedFileError(DomainError):
    code = "UNSUPPORTED_FILE"


class FileTooLargeError(DomainError):
    code = "FILE_TOO_LARGE"


class NoExtractableTextError(DomainError):
    code = "NO_EXTRACTABLE_TEXT"


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    text: str
    kind: str
    metadata: dict[str, int | str]


def _decode_text(content: bytes) -> str:
    if b"\x00" in content:
        raise UnsupportedFileError("Binary content does not match the file extension")
    for encoding in ("utf-8-sig", "utf-8", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnsupportedFileError("Text encoding is not supported")


def parse_document(filename: str, content: bytes) -> ParsedDocument:
    if len(content) > MAX_UPLOAD_BYTES:
        raise FileTooLargeError("Direct uploads must be 4 MB or smaller")
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise UnsupportedFileError("Upload a TXT, Markdown, DOCX, text PDF, or CSV file")
    try:
        if extension in {".txt", ".md"}:
            text = _decode_text(content)
            metadata: dict[str, int | str] = {"bytes": len(content)}
        elif extension == ".docx":
            if not content.startswith(b"PK"):
                raise UnsupportedFileError("File content does not match the DOCX extension")
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                members = archive.infolist()
                if len(members) > MAX_ARCHIVE_FILES:
                    raise UnsupportedFileError("DOCX archive contains too many files")
                if sum(member.file_size for member in members) > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                    raise UnsupportedFileError("DOCX expanded content is too large")
                document_xml = next(
                    (member for member in members if member.filename == "word/document.xml"), None
                )
                if document_xml is None:
                    raise UnsupportedFileError("File content does not match the DOCX extension")
                if document_xml.file_size > MAX_DOCX_XML_BYTES:
                    raise UnsupportedFileError("DOCX document content is too large")
            document = Document(io.BytesIO(content))
            text = "\n\n".join(
                paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()
            )
            metadata = {"paragraphs": len(document.paragraphs)}
        elif extension == ".pdf":
            if not content.lstrip().startswith(b"%PDF-"):
                raise UnsupportedFileError("File content does not match the PDF extension")
            with pymupdf.open(stream=content, filetype="pdf") as pdf:  # type: ignore[no-untyped-call]
                if pdf.needs_pass:
                    raise UnsupportedFileError("Encrypted PDFs are not supported")
                if pdf.page_count > MAX_PDF_PAGES:
                    raise UnsupportedFileError("PDF files may contain at most 200 pages")
                parts: list[str] = []
                character_count = 0
                for page in pdf:
                    part = page.get_text("text")
                    character_count += len(part)
                    if character_count > MAX_EXTRACTED_CHARACTERS:
                        raise FileTooLargeError("Extracted text must be 100,000 characters or fewer")
                    parts.append(part)
                text = "\n\n".join(parts)
                metadata = {"pages": pdf.page_count}
        else:
            decoded = _decode_text(content)
            rows = list(csv.reader(io.StringIO(decoded)))
            if len(rows) > 5_000:
                raise UnsupportedFileError("CSV files may contain at most 5,000 rows")
            text = decoded
            metadata = {"rows": len(rows)}
    except (UnsupportedFileError, FileTooLargeError):
        raise
    except Exception as exc:
        raise UnsupportedFileError("The document is malformed or unsupported") from exc
    if not text.strip():
        raise NoExtractableTextError("This document contains no extractable text")
    if len(text) > MAX_EXTRACTED_CHARACTERS:
        raise FileTooLargeError("Extracted text must be 100,000 characters or fewer")
    return ParsedDocument(text=text, kind=extension.removeprefix("."), metadata=metadata)
